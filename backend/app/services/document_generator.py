import os
import uuid
import html
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from app.schemas.resume import StructuredResume
from app.core.config import settings

def escape_reportlab(text: str) -> str:
    if not text:
        return ""
    escaped = html.escape(text)
    # Restore allowed ReportLab formatting tags
    escaped = escaped.replace("&lt;b&gt;", "<b>").replace("&lt;/b&gt;", "</b>")
    escaped = escaped.replace("&lt;i&gt;", "<i>").replace("&lt;/i&gt;", "</i>")
    escaped = escaped.replace("&lt;br/&gt;", "<br/>").replace("&lt;br&gt;", "<br/>")
    return escaped

def generate_resume_pdf(resume: StructuredResume) -> str:
    filename = f"resume_{uuid.uuid4().hex[:8]}.pdf"
    filepath = os.path.join(settings.EXPORT_DIR, filename)

    doc = SimpleDocTemplate(
        filepath,
        pagesize=letter,
        rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=12
    )
    heading_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=10,
        spaceAfter=4
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6
    )

    story = []

    # Header
    name_str = escape_reportlab(resume.name or "Candidate Resume")
    contact_parts = [escape_reportlab(p) for p in [resume.email, resume.phone, resume.location] if p]
    contact_str = " | ".join(contact_parts)

    story.append(Paragraph(name_str, title_style))
    if contact_str:
        story.append(Paragraph(contact_str, subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CBD5E1"), spaceAfter=10))

    # Summary
    if resume.summary:
        story.append(Paragraph("PROFESSIONAL SUMMARY", heading_style))
        story.append(Paragraph(escape_reportlab(resume.summary), body_style))
        story.append(Spacer(1, 6))

    # Skills
    if resume.skills:
        story.append(Paragraph("SKILLS & TECHNOLOGIES", heading_style))
        skills_str = escape_reportlab(" • ".join(resume.skills))
        story.append(Paragraph(skills_str, body_style))
        story.append(Spacer(1, 6))

    # Experience
    if resume.experience:
        story.append(Paragraph("WORK EXPERIENCE", heading_style))
        for exp in resume.experience:
            header_line = f"<b>{escape_reportlab(exp.title)}</b> — <i>{escape_reportlab(exp.company)}</i> ({escape_reportlab(exp.dates)})"
            story.append(Paragraph(header_line, body_style))
            for bullet in exp.bullets:
                story.append(Paragraph(f"• {escape_reportlab(bullet)}", body_style))
            story.append(Spacer(1, 4))

    # Projects
    if resume.projects:
        story.append(Paragraph("KEY PROJECTS", heading_style))
        for proj in resume.projects:
            story.append(Paragraph(f"<b>{escape_reportlab(proj.name)}</b>: {escape_reportlab(proj.description)}", body_style))
            for bullet in proj.bullets:
                story.append(Paragraph(f"• {escape_reportlab(bullet)}", body_style))
            story.append(Spacer(1, 4))

    # Education
    if resume.education:
        story.append(Paragraph("EDUCATION", heading_style))
        for edu in resume.education:
            story.append(Paragraph(f"<b>{escape_reportlab(edu.degree)} in {escape_reportlab(edu.field)}</b> — {escape_reportlab(edu.institution)} ({escape_reportlab(edu.dates)})", body_style))

    doc.build(story)
    return filepath

def generate_resume_docx(resume: StructuredResume) -> str:
    filename = f"resume_{uuid.uuid4().hex[:8]}.docx"
    filepath = os.path.join(settings.EXPORT_DIR, filename)

    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    name_str = resume.name or "Candidate Resume"
    p_name = doc.add_paragraph()
    run_name = p_name.add_run(name_str)
    run_name.bold = True
    run_name.font.size = Pt(20)
    run_name.font.color.rgb = RGBColor(30, 41, 59)

    contact_parts = [p for p in [resume.email, resume.phone, resume.location] if p]
    if contact_parts:
        p_contact = doc.add_paragraph(" | ".join(contact_parts))
        p_contact.runs[0].font.size = Pt(9.5)
        p_contact.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph("─" * 55)

    if resume.summary:
        h = doc.add_paragraph()
        r = h.add_run("PROFESSIONAL SUMMARY")
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph(resume.summary)

    if resume.skills:
        h = doc.add_paragraph()
        r = h.add_run("SKILLS & TECHNOLOGIES")
        r.bold = True
        r.font.size = Pt(11)
        doc.add_paragraph(" • ".join(resume.skills))

    if resume.experience:
        h = doc.add_paragraph()
        r = h.add_run("WORK EXPERIENCE")
        r.bold = True
        r.font.size = Pt(11)
        for exp in resume.experience:
            p_exp = doc.add_paragraph()
            r_exp = p_exp.add_run(f"{exp.title} | {exp.company} ({exp.dates})")
            r_exp.bold = True
            for b in exp.bullets:
                doc.add_paragraph(f"• {b}")

    if resume.projects:
        h = doc.add_paragraph()
        r = h.add_run("KEY PROJECTS")
        r.bold = True
        r.font.size = Pt(11)
        for proj in resume.projects:
            p_proj = doc.add_paragraph()
            r_proj = p_proj.add_run(f"{proj.name}: {proj.description}")
            r_proj.bold = True
            for b in proj.bullets:
                doc.add_paragraph(f"• {b}")

    if resume.education:
        h = doc.add_paragraph()
        r = h.add_run("EDUCATION")
        r.bold = True
        r.font.size = Pt(11)
        for edu in resume.education:
            doc.add_paragraph(f"{edu.degree} in {edu.field} - {edu.institution} ({edu.dates})")

    doc.save(filepath)
    return filepath

def generate_cover_letter_pdf(content: str, cand_name: str = "Candidate") -> str:
    filename = f"cover_letter_{uuid.uuid4().hex[:8]}.pdf"
    filepath = os.path.join(settings.EXPORT_DIR, filename)

    doc = SimpleDocTemplate(
        filepath,
        pagesize=letter,
        rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CLTitle', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=colors.HexColor('#0F172A'), spaceAfter=14
    )
    body_style = ParagraphStyle(
        'CLBody', parent=styles['Normal'],
        fontName='Helvetica', fontSize=10.5, leading=16, textColor=colors.HexColor('#334155'), spaceAfter=12
    )

    story = [Paragraph(escape_reportlab(f"Cover Letter — {cand_name}"), title_style)]
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CBD5E1"), spaceAfter=14))

    for paragraph in content.split("\n\n"):
        if paragraph.strip():
            story.append(Paragraph(escape_reportlab(paragraph.strip()).replace("\n", "<br/>"), body_style))

    doc.build(story)
    return filepath

def generate_cover_letter_docx(content: str, cand_name: str = "Candidate") -> str:
    filename = f"cover_letter_{uuid.uuid4().hex[:8]}.docx"
    filepath = os.path.join(settings.EXPORT_DIR, filename)

    doc = Document()
    p_title = doc.add_paragraph()
    r = p_title.add_run(f"Cover Letter — {cand_name}")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph("─" * 50)

    for paragraph in content.split("\n\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())

    doc.save(filepath)
    return filepath

