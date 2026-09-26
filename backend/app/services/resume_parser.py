import hashlib
import io
import logging
from typing import Tuple, Optional
from pypdf import PdfReader
from docx import Document
from app.schemas.resume import StructuredResume
from app.ai.base import AIProvider
from app.ai.factory import get_ai_provider
from app.ai.prompts import RESUME_PARSER_SYSTEM

logger = logging.getLogger(__name__)

def compute_file_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()

def extract_text_from_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    text_parts = []
    for page in reader.pages:
        txt = page.extract_text()
        if txt:
            text_parts.append(txt)
    return "\n".join(text_parts).strip()

def extract_text_from_docx(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                text_parts.append(row_text)
    return "\n".join(text_parts).strip()

async def parse_resume_file(
    filename: str,
    content: bytes,
    file_type: str,
    ai_provider: Optional[AIProvider] = None
) -> Tuple[str, StructuredResume, str]:
    file_hash = compute_file_hash(content)
    
    filename_lower = filename.lower()
    if filename_lower.endswith(".pdf") or file_type == "application/pdf":
        raw_text = extract_text_from_pdf(content)
    elif filename_lower.endswith(".docx") or "officedocument" in file_type:
        raw_text = extract_text_from_docx(content)
    elif filename_lower.endswith(".txt") or file_type == "text/plain":
        raw_text = content.decode("utf-8", errors="ignore")
    else:
        try:
            raw_text = extract_text_from_pdf(content)
        except Exception:
            raw_text = content.decode("utf-8", errors="ignore")

    if not raw_text or len(raw_text.strip()) < 10:
        raise ValueError("Could not extract legible text from uploaded file. Please ensure it is not empty or scanned image PDF.")

    provider = ai_provider or get_ai_provider()
    prompt = f"Parse the following resume content into structured format:\n\n{raw_text[:12000]}"
    
    try:
        structured_resume = await provider.generate_structured(
            prompt=prompt,
            schema_cls=StructuredResume,
            system_prompt=RESUME_PARSER_SYSTEM
        )
    except Exception as e:
        logger.warning(f"AI parsing failed, creating fallback structured resume: {e}")
        structured_resume = StructuredResume(
            summary=raw_text[:300],
            skills=["Extracted from document"],
            experience=[]
        )

    return raw_text, structured_resume, file_hash

