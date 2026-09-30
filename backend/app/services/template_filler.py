"""Apply accepted edits onto the user's ORIGINAL uploaded file so the download keeps their template."""
import difflib
import html
import os
import re
import uuid
from typing import List, Optional, Tuple

import pymupdf
from docx import Document

from app.core.config import settings

Pair = Tuple[str, str]


def _norm(s: str) -> str:
    return " ".join((s or "").split())


def _skills_pair(lines: List[str], skills: List[str], new_skills: List[str]) -> Optional[Pair]:
    """Find the line listing the most existing skills and append new ones with its own separator."""
    new_skills = [s for s in new_skills if s]
    if not new_skills or not skills:
        return None
    best = max(lines, key=lambda l: sum(s.lower() in l.lower() for s in skills), default="")
    if sum(s.lower() in best.lower() for s in skills) < 2:
        return None
    sep = next((s for s in (" • ", " | ", ", ", " · ") if s in best), ", ")
    return best, best.rstrip() + sep + sep.join(new_skills)


# ---------------- DOCX ----------------

def _docx_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _replace_in_paragraph(p, before: str, after: str) -> bool:
    # Exact hit inside one run: formatting untouched.
    for run in p.runs:
        if before in run.text:
            run.text = run.text.replace(before, after, 1)
            return True
    text = p.text
    if _norm(before) in _norm(text):
        new_text = _norm(text).replace(_norm(before), after, 1)
    elif difflib.SequenceMatcher(None, _norm(before), _norm(text)).ratio() >= 0.85:
        prefix = re.match(r"^[\W_]*", text).group(0)  # keep literal bullet chars like "• "
        new_text = prefix + after
    else:
        return False
    # ponytail: collapses mixed formatting inside this one paragraph to its first run's style
    runs = p.runs
    if not runs:
        return False
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""
    return True


def fill_docx(src_path: str, pairs: List[Pair], skills: List[str], new_skills: List[str]) -> str:
    doc = Document(src_path)
    paras = list(_docx_paragraphs(doc))
    sp = _skills_pair([p.text for p in paras], skills, new_skills)
    for before, after in pairs + ([sp] if sp else []):
        for p in paras:
            if _replace_in_paragraph(p, before, after):
                break
    out = os.path.join(settings.EXPORT_DIR, f"resume_{uuid.uuid4().hex[:8]}.docx")
    doc.save(out)
    return out


# ---------------- PDF ----------------

def _pdf_rects(page, text: str) -> List[pymupdf.Rect]:
    text = _norm(text)
    hits = page.search_for(text)
    if hits:
        return hits
    # AI-parsed text may differ slightly from the PDF: anchor on first/last words instead.
    words = text.split()
    if len(words) < 6:
        return []
    heads = page.search_for(" ".join(words[:5]))
    tails = page.search_for(" ".join(words[-5:]))
    if not heads or not tails:
        return []
    s = heads[0]
    e = next((t for t in tails if t.y0 >= s.y0 - 1), None)
    if e is None:
        return []
    rects = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            r = pymupdf.Rect(line["bbox"])
            if r.y1 > s.y0 + 1 and r.y0 < e.y1 - 1 and (r.intersects(s) or abs(r.x0 - s.x0) < 30):
                rects.append(r)
    if rects:
        rects[0].x0 = max(rects[0].x0, s.x0)  # don't wipe the bullet glyph
    return rects


def _style_at(page, rect) -> Tuple[float, str, bool]:
    for block in page.get_text("dict", clip=rect)["blocks"]:
        for line in block.get("lines", []):
            for span in line["spans"]:
                if span["text"].strip():
                    return span["size"], f"#{span['color']:06x}", bool(span["flags"] & 16)
    return 10.0, "#000000", False


def _text_box(page, rects) -> pymupdf.Rect:
    box = pymupdf.Rect(rects[0])
    for r in rects[1:]:
        box |= r
    box.x0 = rects[0].x0  # continuation lines start under the bullet; keep text right of it
    # Widen to the column's right edge so longer text wraps like the original instead of shrinking.
    box.x1 = max([box.x1] + [b[2] for b in page.get_text("blocks") if abs(b[0] - box.x0) < 30])
    return box


def fill_pdf(src_path: str, pairs: List[Pair], skills: List[str], new_skills: List[str]) -> str:
    doc = pymupdf.open(src_path)
    lines = [l for page in doc for l in page.get_text().splitlines()]
    sp = _skills_pair(lines, skills, new_skills)
    # Locate + read styles on the untouched page first; redacting in between would erase earlier inserts.
    edits = []
    for before, after in pairs + ([sp] if sp else []):
        for page in doc:
            rects = _pdf_rects(page, before)
            if rects:
                edits.append((page, rects, after, _style_at(page, rects[0]), _text_box(page, rects)))
                break
    for page, rects, *_ in edits:
        for r in rects:
            h = r.height * 0.3  # adjacent line boxes overlap; only hit the middle band
            page.add_redact_annot(pymupdf.Rect(r.x0, r.y0 + h, r.x1, r.y1 - h), fill=False)  # transparent: keeps backgrounds
    for page in {id(e[0]): e[0] for e in edits}.values():
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
    for page, _, after, (size, color, bold), box in edits:
        css = (f"* {{font-family: sans-serif; font-size: {size:.1f}px; line-height: 1.15; color: {color}; "
               f"font-weight: {'bold' if bold else 'normal'}; margin: 0;}}")
        # ponytail: original embedded font is usually subset, so we fall back to sans-serif at same size/colour;
        # longer text is auto-shrunk to fit the old box (scale_low=0).
        page.insert_htmlbox(box, html.escape(after), css=css, scale_low=0)
    out = os.path.join(settings.EXPORT_DIR, f"resume_{uuid.uuid4().hex[:8]}.pdf")
    doc.save(out, garbage=3, deflate=True)
    doc.close()
    return out


def fill_original_template(src_path: str, want_ext: str, pairs: List[Pair], skills: List[str], new_skills: List[str]) -> Optional[str]:
    """Returns path of edited copy, or None if the original can't be used for this format."""
    if not src_path or not os.path.exists(src_path) or os.path.splitext(src_path)[1].lower() != want_ext:
        return None
    return (fill_pdf if want_ext == ".pdf" else fill_docx)(src_path, pairs, skills, new_skills)
