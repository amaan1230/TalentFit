import pymupdf
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph
from app.services.template_filler import fill_original_template

LONG = "Built REST APIs in Python and FastAPI serving internal dashboards for the finance team across three regions"
PAIRS = [("Worked on backend stuff", "Engineered scalable FastAPI backend services"), (LONG, "Architected Python/FastAPI REST APIs powering finance dashboards in three regions")]
SKILLS, NEW = ["Python", "FastAPI", "Docker"], ["MongoDB"]


def test_pdf(tmp_path):
    src = str(tmp_path / "cv.pdf")
    st = getSampleStyleSheet()
    SimpleDocTemplate(src, pagesize=letter).build([Paragraph("<font color='#1d4ed8'>Jane Doe</font>", st["Title"]),
        Paragraph("Python, FastAPI, Docker", st["Normal"]), Paragraph("• Worked on backend stuff", st["Normal"]), Paragraph("• " + LONG, st["Normal"])])
    out = fill_original_template(src, ".pdf", PAIRS, SKILLS, NEW)
    text = " ".join(pymupdf.open(out)[0].get_text().split())
    assert "Engineered scalable FastAPI backend services" in text and "Worked on backend" not in text
    assert "Architected" in text and "internal dashboards" not in text
    assert "MongoDB" in text and "Jane Doe" in text
    pymupdf.open(out)[0].get_pixmap(dpi=90).save(str(tmp_path / "out.png"))
    print(tmp_path / "out.png")


def test_docx(tmp_path):
    src = str(tmp_path / "cv.docx")
    d = Document(); d.add_heading("Jane Doe"); d.add_paragraph("Python, FastAPI, Docker")
    r = d.add_paragraph("Worked on backend stuff", style="List Bullet").runs[0]; r.bold = True
    d.add_paragraph(LONG, style="List Bullet"); d.save(src)
    out = Document(fill_original_template(src, ".docx", PAIRS, SKILLS, NEW))
    texts = [p.text for p in out.paragraphs]
    assert "Engineered scalable FastAPI backend services" in texts and out.paragraphs[2].runs[0].bold
    assert out.paragraphs[2].style.name == "List Bullet"
    assert "Python, FastAPI, Docker, MongoDB" in texts and texts[3].startswith("Architected")


def test_format_mismatch_falls_back(tmp_path):
    assert fill_original_template(str(tmp_path / "x.pdf"), ".docx", [], [], []) is None
